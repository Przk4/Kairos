# tu_app/rag_service.py
"""
Canvas RAG service — downloads course files, extracts text, and indexes them
via the generic VectorStore module.

All embedding / storage / search operations are delegated to VectorStore.
This file only contains Canvas-specific logic:
  - Document extraction (PDF, PPTX, DOCX, XLSX, images via OCR)
  - Canvas file download (Bearer-token auth, ngrok fallback)
  - Slide-noise cleaning
  - Chunk-importance heuristics (identifies definitions, examples, position, etc.)
  - analyze_module()  — main ingestion pipeline
  - search_context()  — thin wrapper calling VectorStore.search()
"""

import os
import logging
import warnings
import re
from pathlib import Path
from typing import List, Optional

import requests
from django.conf import settings

from tu_app.debug_service import get_debug_service
from tu_app.vector_store import VectorStore

logger = logging.getLogger(__name__)

# Suppress noisy Pydantic v1 warnings from chromadb
warnings.filterwarnings('ignore', message='Core Pydantic V1 functionality')
warnings.filterwarnings('ignore', category=UserWarning)

# ---------------------------------------------------------------------------
# Optional document-format libraries
# ---------------------------------------------------------------------------
try:
    import pdfplumber
    PDF_READER_AVAILABLE = True
    logger.info("pdfplumber imported successfully")
except ImportError:
    logger.warning("pdfplumber not available — PDF processing will be limited")
    PDF_READER_AVAILABLE = False
    pdfplumber = None

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
    logger.info("python-pptx imported successfully")
except ImportError:
    logger.warning("python-pptx not available — PPTX processing disabled")
    PPTX_AVAILABLE = False
    Presentation = None

try:
    from docx import Document
    DOCX_AVAILABLE = True
    logger.info("python-docx imported successfully")
except ImportError:
    logger.warning("python-docx not available — DOCX processing disabled")
    DOCX_AVAILABLE = False
    Document = None

try:
    from openpyxl import load_workbook
    XLSX_AVAILABLE = True
    logger.info("openpyxl imported successfully")
except ImportError:
    logger.warning("openpyxl not available — XLSX processing disabled")
    XLSX_AVAILABLE = False
    load_workbook = None


class RAGService:
    """
    Canvas RAG service.

    Delegates all vector-store operations (embed, store, search) to
    :class:`tu_app.vector_store.VectorStore`.  Only Canvas-specific logic
    lives here.
    """

    def __init__(self, persistent_dir: Optional[str] = None):
        persist_dir = persistent_dir or os.path.join(settings.BASE_DIR, '.chroma_db')
        json_dir = os.path.join(settings.BASE_DIR, '.embeddings')
        self.vector_store = VectorStore(persist_dir=persist_dir, json_dir=json_dir)
        logger.info(
            f"RAG Service initialised — ChromaDB: {self.vector_store.using_chromadb}, "
            f"PDF: {PDF_READER_AVAILABLE}"
        )

    # ------------------------------------------------------------------
    # Compatibility properties — views.py accesses these directly
    # ------------------------------------------------------------------

    @property
    def using_chromadb(self) -> bool:
        return self.vector_store.using_chromadb

    @property
    def client(self):
        return self.vector_store.client

    @property
    def persistent_dir(self) -> str:
        return self.vector_store.persist_dir

    @property
    def in_memory_db(self) -> dict:
        return self.vector_store._memory

    @property
    def embedding_model(self):
        return self.vector_store.embedding_model

    @property
    def json_embeddings_dir(self) -> str:
        return self.vector_store.json_dir

    # ------------------------------------------------------------------
    # PDF extraction
    # ------------------------------------------------------------------

    def _extract_text_from_pdf(self, pdf_bytes) -> str:
        if not PDF_READER_AVAILABLE:
            logger.warning("pdfplumber not available, returning placeholder")
            return "[PDF Document - Text extraction not available]"
        try:
            import io
            logger.info(f"Extracting PDF ({len(pdf_bytes)} bytes) with pdfplumber")
            with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                logger.info(f"PDF pages: {len(pdf.pages)}")
                text = ""
                for page_num, page in enumerate(pdf.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text += f"\n--- Página {page_num + 1} ---\n{page_text}"
                    else:
                        logger.warning(f"Page {page_num + 1}: no text extracted")
                return text if text.strip() else "[PDF vacío o no legible]"
        except Exception as e:
            logger.error(f"Error extracting PDF text: {e}", exc_info=True)
            return f"[Error leyendo PDF: {str(e)[:100]}]"

    # ------------------------------------------------------------------
    # PPTX helpers
    # ------------------------------------------------------------------

    _NOISE_PATTERNS = [
        re.compile(r'^\d{1,3}$'),
        re.compile(r'^20[\dXx]{2}$'),
        re.compile(r'pie de p[aá]gina', re.IGNORECASE),
        re.compile(r'ejemplo de texto', re.IGNORECASE),
        re.compile(r'^(?:©|\(c\)|copyright)\s*20', re.IGNORECASE),
        re.compile(r'^footer\b', re.IGNORECASE),
        re.compile(r'^(click to edit|haga clic|insert|placeholder)', re.IGNORECASE),
        re.compile(r'^(slide|diapositiva)\s*\d', re.IGNORECASE),
        re.compile(r'^\*+$'),
    ]

    def _clean_slide_text(self, text: str) -> str:
        if not text:
            return ""
        lines = []
        for line in text.split('\n'):
            line = line.strip()
            if len(line) < 3:
                continue
            if any(p.search(line) for p in self._NOISE_PATTERNS):
                continue
            lines.append(line)
        return '\n'.join(lines)

    def _table_to_markdown(self, table) -> str:
        rows_data = []
        for row in table.rows:
            cells = [cell.text.strip().replace('|', '/') for cell in row.cells]
            rows_data.append(cells)
        if not rows_data:
            return ""
        max_cols = max(len(r) for r in rows_data)
        for r in rows_data:
            while len(r) < max_cols:
                r.append('')
        md_lines = [
            '| ' + ' | '.join(rows_data[0]) + ' |',
            '| ' + ' | '.join(['---'] * max_cols) + ' |',
        ]
        for row in rows_data[1:]:
            md_lines.append('| ' + ' | '.join(row) + ' |')
        return '\n'.join(md_lines)

    def _extract_text_from_pptx(self, pptx_bytes) -> Optional[str]:
        if not PPTX_AVAILABLE:
            logger.warning("python-pptx not available")
            return None
        try:
            import io
            logger.info(f"Extracting PPTX ({len(pptx_bytes)} bytes)")
            presentation = Presentation(io.BytesIO(pptx_bytes))
            logger.info(f"PPTX slides: {len(presentation.slides)}")
            all_slide_texts = []
            skipped = 0
            self._last_pptx_tables = []

            for slide_num, slide in enumerate(presentation.slides):
                slide_title = ""
                body_parts = []

                for shape in slide.shapes:
                    try:
                        shape_text = ""
                        try:
                            if hasattr(shape, 'text') and shape.text:
                                shape_text = shape.text.strip()
                        except Exception:
                            pass

                        is_title = False
                        try:
                            if (hasattr(shape, 'placeholder_format')
                                    and shape.placeholder_format is not None
                                    and shape.placeholder_format.idx in (0, 1)):
                                is_title = True
                        except Exception:
                            pass

                        is_table = False
                        try:
                            if shape.has_table:
                                is_table = True
                                md_table = self._table_to_markdown(shape.table)
                                if md_table and len(md_table) > 10:
                                    body_parts.append(md_table)
                                    self._last_pptx_tables.append({
                                        'slide_num': slide_num + 1,
                                        'slide_title': slide_title or f"Slide {slide_num + 1}",
                                        'markdown': md_table,
                                        'row_count': len(shape.table.rows),
                                        'col_count': len(shape.table.columns) if hasattr(shape.table, 'columns') else 0,
                                    })
                        except Exception:
                            pass

                        if shape_text and not is_table:
                            if is_title and not slide_title:
                                slide_title = shape_text
                            else:
                                body_parts.append(shape_text)
                    except Exception as e:
                        logger.warning(f"Error in shape (slide {slide_num + 1}): {e}")
                        continue

                raw_body = '\n'.join(body_parts)
                clean_body = self._clean_slide_text(raw_body)
                clean_title = self._clean_slide_text(slide_title)

                total_content = f"{clean_title} {clean_body}".strip()
                if len(total_content) < 10:
                    skipped += 1
                    continue

                block = (
                    f"\n{clean_title}\n{clean_body}" if clean_body
                    else f"\n{clean_title}"
                ) if clean_title else f"\n{clean_body}"
                all_slide_texts.append(block)

            if not all_slide_texts:
                return "[PPTX vacío o sin contenido de texto]"

            final_text = '\n'.join(all_slide_texts)
            logger.info(
                f"PPTX: {len(all_slide_texts)} slides kept, {skipped} skipped, "
                f"{len(self._last_pptx_tables)} tables, {len(final_text)} chars"
            )
            return final_text
        except Exception as e:
            logger.error(f"Error extracting PPTX: {e}", exc_info=True)
            return None

    def _extract_text_from_docx(self, docx_bytes) -> Optional[str]:
        if not DOCX_AVAILABLE:
            logger.warning("python-docx not available")
            return None
        try:
            import io
            document = Document(io.BytesIO(docx_bytes))
            text = ""
            for paragraph in document.paragraphs:
                para_text = paragraph.text.strip()
                if para_text:
                    text += para_text + "\n"
            for table_num, table in enumerate(document.tables):
                text += f"\n--- Tabla {table_num + 1} ---\n"
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text += " | ".join(row_text) + "\n"
            return text if text.strip() else "[DOCX vacío o sin contenido]"
        except Exception as e:
            logger.error(f"Error extracting DOCX: {e}", exc_info=True)
            return None

    def _extract_text_from_xlsx(self, xlsx_bytes) -> Optional[str]:
        if not XLSX_AVAILABLE:
            logger.warning("openpyxl not available")
            return None
        try:
            import io
            workbook = load_workbook(io.BytesIO(xlsx_bytes))
            text = ""
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                text += f"\n--- Sheet: {sheet_name} ---\n"
                for row in worksheet.iter_rows(values_only=True):
                    row_text = [str(v).strip() for v in row if v is not None]
                    if row_text:
                        text += " | ".join(row_text) + "\n"
            return text if text.strip() else "[XLSX vacío o sin contenido]"
        except Exception as e:
            logger.error(f"Error extracting XLSX: {e}", exc_info=True)
            return None

    # ------------------------------------------------------------------
    # Canvas file download + routing
    # ------------------------------------------------------------------

    def _extract_text_from_file(self, file_url: str, access_token: str) -> Optional[str]:
        """Download a Canvas file (with Bearer auth) and extract its text."""
        try:
            logger.info(f"[FILE DOWNLOAD] {file_url}")
            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36',
                'Accept': '*/*',
                'Accept-Encoding': 'gzip, deflate',
            }
            if access_token:
                headers['Authorization'] = f'Bearer {access_token}'
            else:
                logger.warning("[FILE DOWNLOAD] No access token provided")

            url_to_request = file_url
            if '/api/v1/files/' in file_url and '/download' not in file_url:
                url_to_request = f"{file_url}/download"

            response = requests.get(
                url_to_request,
                headers=headers,
                timeout=60,
                allow_redirects=True,
                verify=False,  # ngrok SSL
            )
            logger.info(
                f"[FILE DOWNLOAD] {response.status_code}, {len(response.content)} bytes, "
                f"Content-Type: {response.headers.get('content-type', '?')}"
            )

            if response.status_code != 200:
                logger.error(f"[FILE DOWNLOAD] HTTP {response.status_code}")
                return None

            file_content = response.content

            # If we got HTML back (ngrok redirect), try localhost
            if file_content.startswith(b'<!DOCTYPE') or file_content.startswith(b'<html'):
                logger.warning("[FILE DOWNLOAD] Got HTML — retrying with localhost")
                if 'ngrok' in url_to_request:
                    localhost_url = url_to_request.replace(
                        url_to_request.split('/courses')[0], 'http://127.0.0.1:3000'
                    )
                    response = requests.get(localhost_url, headers=headers, timeout=60,
                                            allow_redirects=True, verify=False)
                    file_content = response.content
                    if file_content.startswith(b'<!DOCTYPE') or file_content.startswith(b'<html'):
                        logger.error("[FILE DOWNLOAD] Still HTML on localhost — giving up")
                        return None
                else:
                    logger.error("[FILE DOWNLOAD] Got HTML instead of file")
                    return None

            # Cache raw bytes for downstream image extraction (OCR)
            self._last_file_bytes = file_content

            file_lower = file_url.lower()
            content_type = response.headers.get('content-type', '').lower()

            # --- Route by extension / content-type ---
            if file_lower.endswith('.pdf') or 'application/pdf' in content_type:
                text = self._extract_text_from_pdf(file_content)
                return text if text and text.strip() and not text.startswith('[') else None

            elif file_lower.endswith(('.pptx', '.ppt')) or 'presentation' in content_type:
                text = self._extract_text_from_pptx(file_content)
                return text if text and text.strip() else None

            elif (file_lower.endswith(('.docx', '.doc'))
                  or 'wordprocessingml' in content_type or 'msword' in content_type):
                text = self._extract_text_from_docx(file_content)
                return text if text and text.strip() else None

            elif (file_lower.endswith(('.xlsx', '.xls'))
                  or 'spreadsheetml' in content_type or 'ms-excel' in content_type):
                text = self._extract_text_from_xlsx(file_content)
                return text if text and text.strip() else None

            elif file_lower.endswith('.txt') or 'text/plain' in content_type:
                return file_content.decode('utf-8', errors='ignore')

            elif (file_lower.endswith(('.jpg', '.jpeg', '.png', '.webp', '.bmp', '.gif', '.tiff'))
                  or content_type.startswith('image/')):
                try:
                    from .ocr_service import get_ocr_service
                    ocr = get_ocr_service()
                    text = ocr.extract_text(file_content, detail="high")
                    return text if text and text.strip() else None
                except Exception as e:
                    logger.error(f"[IMAGE OCR] {e}")
                    return None

            else:
                # Magic-byte detection
                if file_content.startswith(b'%PDF'):
                    return self._extract_text_from_pdf(file_content)

                if file_content.startswith(b'PK\x03\x04'):
                    try:
                        import zipfile, io as _io
                        with zipfile.ZipFile(_io.BytesIO(file_content)) as zf:
                            names = zf.namelist()
                            if any('slide' in n.lower() for n in names):
                                return self._extract_text_from_pptx(file_content)
                            elif 'word/document.xml' in names:
                                return self._extract_text_from_docx(file_content)
                            elif 'xl/workbook.xml' in names:
                                return self._extract_text_from_xlsx(file_content)
                    except Exception:
                        pass

                # Try plain text
                try:
                    decoded = file_content.decode('utf-8', errors='strict')
                    if decoded.strip():
                        return decoded
                except UnicodeDecodeError:
                    pass

                # Try image via OCR (JPEG / PNG / WEBP magic bytes)
                if (file_content[:3] == b'\xff\xd8\xff'
                        or file_content[:4] == b'\x89PNG'
                        or (file_content[:4] == b'RIFF' and file_content[8:12] == b'WEBP')):
                    try:
                        from .ocr_service import get_ocr_service
                        ocr = get_ocr_service()
                        text = ocr.extract_text(file_content, detail="high")
                        return text if text and text.strip() else None
                    except Exception as e:
                        logger.warning(f"[UNKNOWN OCR] {e}")

                logger.warning("[FILE DOWNLOAD] Cannot determine file type or extract text")
                return None

        except requests.exceptions.RequestException as e:
            logger.error(f"[FILE DOWNLOAD] {type(e).__name__}: {e}")
            return None
        except Exception as e:
            logger.error(f"[FILE PROCESSING] {type(e).__name__}: {e}")
            return None

    # ------------------------------------------------------------------
    # Chunk-importance heuristics  (Canvas-specific)
    # ------------------------------------------------------------------

    def _calculate_chunk_importance(
        self,
        chunk: str,
        chunk_index: int,
        total_chunks: int,
        full_text: str = None,
        document_title: str = None,
    ) -> float:
        """
        Heuristic importance score (0.1 – 1.0) for a text chunk.
        Higher = more likely to be a central concept rather than an example/detail.
        """
        if not chunk or not chunk.strip():
            return 0.3

        chunk_lower = chunk.lower()
        score = 0.5

        # --- Definition / core-concept signals (+) ---
        for pattern in [
            r'\b(es|son|se define como|significa|consiste en|se entiende por)\b',
            r'\b(definici[oó]n|concepto|principio|ley|teor[ií]a|fundamento)\b',
            r'\b(importante|esencial|fundamental|clave|b[aá]sico|principal)\b',
            r'\b(objetivo|prop[oó]sito|meta|fin|conclusi[oó]n)\b',
        ]:
            if re.search(pattern, chunk_lower):
                score += 0.08

        # --- Example / detail signals (-) ---
        for pattern in [
            r'\b(por ejemplo|ejemplo|e\.g\.|i\.e\.|como muestra|ilustra)\b',
            r'\b(caso particular|en este caso|supongamos|imagina)\b',
            r'\b(espec[ií]ficamente|particularmente|en concreto)\b',
            r'\b(ver figura|ver tabla|como se observa|seg[uú]n la imagen)\b',
        ]:
            if re.search(pattern, chunk_lower):
                score -= 0.05

        # --- Position ---
        if total_chunks > 1:
            position_ratio = chunk_index / (total_chunks - 1)
            if chunk_index == 0:
                score += 0.12
            elif chunk_index == total_chunks - 1:
                score += 0.10
            elif position_ratio < 0.25:
                score += 0.05
            elif position_ratio > 0.75:
                score += 0.03

        # --- Structure / formatting ---
        for pattern, bonus in [
            (r'^[A-Z][A-Z\s]+$', 0.15),
            (r'^\d+[\.\)]\s', 0.08),
            (r'^[-•]\s', 0.05),
            (r'\b(paso \d|etapa \d|fase \d)\b', 0.06),
            (r'\b(primero|segundo|tercero|finalmente)\b', 0.04),
        ]:
            if re.search(pattern, chunk_lower, re.MULTILINE):
                score += bonus

        # --- Information density ---
        words = chunk_lower.split()
        if len(words) > 10:
            filler = {'el', 'la', 'los', 'las', 'de', 'del', 'al', 'a', 'en', 'con',
                      'por', 'para', 'que', 'y', 'o', 'un', 'una', 'como', 'es', 'son',
                      'se', 'su', 'sus', 'este', 'esta', 'esto'}
            content_ratio = 1 - sum(1 for w in words if w in filler) / len(words)
            if content_ratio > 0.7:
                score += 0.08
            elif content_ratio > 0.6:
                score += 0.04

        # --- Chunk length ---
        if 200 < len(chunk) < 1200:
            score += 0.05
        elif len(chunk) < 100:
            score -= 0.05

        # --- Document title overlap ---
        if document_title:
            title_words = set(re.findall(r'\b[a-záéíóúñü]{4,}\b', document_title.lower()))
            chunk_words = set(re.findall(r'\b[a-záéíóúñü]{4,}\b', chunk_lower))
            overlap = len(title_words & chunk_words)
            if overlap > 0:
                score += min(0.10, overlap * 0.03)

        return max(0.1, min(1.0, score))

    def _calculate_importance_batch(
        self,
        chunks: List[str],
        document_title: str = None,
    ) -> List[float]:
        if not chunks:
            return []
        full_text = '\n'.join(chunks)
        total = len(chunks)
        scores = [
            self._calculate_chunk_importance(c, i, total, full_text, document_title)
            for i, c in enumerate(chunks)
        ]
        if scores:
            logger.debug(
                f"[IMPORTANCE] {len(scores)} chunks, avg={sum(scores)/len(scores):.3f}, "
                f"min={min(scores):.3f}, max={max(scores):.3f}"
            )
        return scores

    # ------------------------------------------------------------------
    # Canvas ingestion pipeline
    # ------------------------------------------------------------------

    def analyze_module(self, module, user_token: str) -> bool:
        """
        Process every ModuleItem in *module*:
        download file → extract text → semantic chunk → compute importance → upsert.
        """
        debug_service = get_debug_service()
        try:
            collection_name = f"module_{module.id}_course_{module.course.id}"
            debug_service.log_rag_analysis_start(module.id, module.name)
            logger.info(f"Starting module analysis: {collection_name}")

            # Clear stale data
            self.vector_store.delete_collection(collection_name)

            items = module.items.all()
            embeddings_count = 0
            analyzed_documents = []

            for i, item in enumerate(items):
                try:
                    logger.info(f"[ITEM {i+1}/{len(items)}] {item.title}")

                    text = f"{item.title} - {item.get_item_type_display()}"
                    extracted = None
                    is_external_content = False

                    if item.url:
                        extracted = self._extract_text_from_file(item.url, user_token)
                        if extracted:
                            text = f"{text}\n\n{extracted}"
                            is_external_content = True
                        else:
                            logger.warning(f"  [EXTRACT] Failed — using title only: '{text}'")

                    # Reject binary PDF that leaked through
                    if is_external_content and text.startswith('%PDF'):
                        logger.error("  [VALIDATE] Raw PDF binary — skipping item")
                        continue

                    # Accumulate per-item docs for a single upsert call
                    docs, metas, ids, embs = [], [], [], []

                    if extracted:
                        raw_chunks = self.vector_store.chunk_semantic(
                            text, max_chunk_size=1000, similarity_threshold=0.45
                        )
                        chunks = [rc for rc in raw_chunks if rc and len(rc.strip()) > 15]
                        if not chunks:
                            chunks = [text.split('\n')[0]]

                        logger.info(f"  [CHUNKS] {len(text)} chars → {len(raw_chunks)} raw → {len(chunks)} clean")

                        importance_scores = self._calculate_importance_batch(
                            chunks, document_title=item.title
                        )

                        for chunk_idx, chunk in enumerate(chunks):
                            importance = (
                                importance_scores[chunk_idx]
                                if chunk_idx < len(importance_scores) else 0.5
                            )
                            ids.append(f"item_{item.id}_chunk_{chunk_idx}")
                            docs.append(chunk)
                            embs.append(self.vector_store.embedding_model.encode(chunk).tolist())
                            metas.append({
                                'item_id': str(item.id),
                                'item_title': item.title,
                                'item_type': item.get_item_type_display(),
                                'module_id': str(module.id),
                                'chunk_index': str(chunk_idx),
                                'chunk_total': str(len(chunks)),
                                'has_content': 'yes',
                                'importance_score': str(round(importance, 3)),
                            })

                        # Separate embeddings for PPTX tables
                        pptx_tables = getattr(self, '_last_pptx_tables', [])
                        for tbl_idx, tbl in enumerate(pptx_tables):
                            table_text = (
                                f"{item.title} - Tabla en: {tbl['slide_title']}\n\n{tbl['markdown']}"
                            )
                            table_text_clean = self._clean_slide_text(table_text)
                            if not table_text_clean or len(table_text_clean.strip()) < 20:
                                continue
                            ids.append(f"item_{item.id}_table_{tbl_idx}")
                            docs.append(table_text_clean)
                            embs.append(
                                self.vector_store.embedding_model.encode(table_text_clean).tolist()
                            )
                            metas.append({
                                'item_id': str(item.id),
                                'item_title': item.title,
                                'item_type': item.get_item_type_display(),
                                'module_id': str(module.id),
                                'content_type': 'table',
                                'table_slide': str(tbl['slide_num']),
                                'table_slide_title': tbl['slide_title'],
                                'table_rows': str(tbl['row_count']),
                                'has_content': 'yes',
                                'importance_score': '0.75',
                            })
                        self._last_pptx_tables = []

                        # OCR embedded images (PPTX / PDF)
                        if item.url:
                            try:
                                from .ocr_service import get_ocr_service
                                ocr = get_ocr_service()
                                file_lower = item.url.lower()
                                ocr_images = []
                                raw_bytes = getattr(self, '_last_file_bytes', None)
                                if raw_bytes:
                                    if file_lower.endswith(('.pptx', '.ppt')):
                                        ocr_images = ocr.extract_images_from_pptx(raw_bytes)
                                    elif file_lower.endswith('.pdf'):
                                        ocr_images = ocr.extract_images_from_pdf(raw_bytes)
                                for img_idx, img_data in enumerate(ocr_images):
                                    img_text = img_data.get('text', '')
                                    if not img_text or len(img_text.strip()) < 20:
                                        continue
                                    location = img_data.get('slide', img_data.get('page', '?'))
                                    img_doc = (
                                        f"{item.title} - Imagen (slide/página {location})"
                                        f"\n\n{img_text}"
                                    )
                                    ids.append(f"item_{item.id}_img_{img_idx}")
                                    docs.append(img_doc)
                                    embs.append(
                                        self.vector_store.embedding_model.encode(img_doc).tolist()
                                    )
                                    metas.append({
                                        'item_id': str(item.id),
                                        'item_title': item.title,
                                        'item_type': item.get_item_type_display(),
                                        'module_id': str(module.id),
                                        'content_type': 'image_ocr',
                                        'image_location': str(location),
                                        'has_content': 'yes',
                                        'importance_score': '0.7',
                                    })
                                if ocr_images:
                                    logger.info(f"  [IMG OCR] {len(ocr_images)} image embeddings")
                            except Exception as e:
                                logger.warning(f"  [IMG OCR] skipped: {e}")

                    else:
                        # Title-only fallback
                        ids.append(f"item_{item.id}")
                        docs.append(text)
                        embs.append(self.vector_store.embedding_model.encode(text).tolist())
                        metas.append({
                            'item_id': str(item.id),
                            'item_title': item.title,
                            'item_type': item.get_item_type_display(),
                            'module_id': str(module.id),
                            'has_content': 'no',
                            'importance_score': '0.3',
                        })

                    if docs:
                        self.vector_store.upsert(collection_name, docs, metas, ids, embs)
                        embeddings_count += len(docs)
                        debug_service.log_embedding_created(item.id, item.title, len(embs[-1]))

                except Exception as e:
                    logger.error(f"Error processing item {item.id}: {e}")
                    debug_service.log_error("EMBEDDING", f"Error procesando item {item.id}", e)
                    continue

            debug_service.log_analyzed_documents(module.id, module.name, analyzed_documents)
            debug_service.log_rag_analysis_complete(module.id, len(items), embeddings_count)
            logger.info(f"Module analysis complete: {len(items)} items, {embeddings_count} embeddings")

            if embeddings_count == 0:
                logger.error(f"[ANALYZE] 0 embeddings for module {module.id} — analysis failed")
                return False

            return True

        except Exception as e:
            logger.error(f"Error analyzing module: {e}")
            debug_service.log_error("RAG_ANALYSIS", f"Error en analyze_module módulo {module.id}", e)
            return False

    # ------------------------------------------------------------------
    # Search — thin wrapper around VectorStore
    # ------------------------------------------------------------------

    def search_context(
        self,
        query: str,
        module_id: int,
        course_id: int,
        top_k: int = 5,
        query_type: str = 'general',
        importance_weight: Optional[float] = None,
    ) -> List[dict]:
        """
        Semantic search over a Canvas module collection.

        Returns a list of dicts with keys:
            content, metadata, relevance_score, importance_score, combined_score.
        """
        debug_service = get_debug_service()
        try:
            collection_name = f"module_{module_id}_course_{course_id}"
            results = self.vector_store.search(
                collection=collection_name,
                query=query,
                top_k=top_k,
                query_type=query_type,
                importance_weight=importance_weight,
            )
            debug_service.log_rag_search(query, module_id, len(results))
            return results
        except Exception as e:
            logger.error(f"Error in search_context: {e}")
            debug_service.log_error("SEARCH", "Error en search_context", e)
            return []

    # ------------------------------------------------------------------
    # Collection management — thin wrappers
    # ------------------------------------------------------------------

    def delete_collection(self, module_id: int, course_id: int) -> None:
        collection_name = f"module_{module_id}_course_{course_id}"
        self.vector_store.delete_collection(collection_name)

    def check_collection_exists(self, collection_name: str) -> bool:
        return self.vector_store.collection_exists(collection_name)

    def get_embeddings_count(self, collection_name: str) -> int:
        return self.vector_store.collection_count(collection_name)


# ---------------------------------------------------------------------------
# Singleton accessor
# ---------------------------------------------------------------------------
_rag_service: Optional[RAGService] = None


def get_rag_service() -> RAGService:
    """Return the shared RAGService instance (lazy-initialised)."""
    global _rag_service
    if _rag_service is None:
        _rag_service = RAGService()
    return _rag_service
